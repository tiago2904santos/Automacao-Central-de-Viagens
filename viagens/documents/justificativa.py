# viagens/documents/justificativa.py
"""Geração de documento de Justificativa a partir do modelo DOCX com placeholders {{chave}}."""
from __future__ import annotations

import re
from datetime import date
from io import BytesIO
from pathlib import Path

from django.conf import settings

from .document import (
    _build_endereco_formatado,
    _extract_placeholders,
    _iter_docx_xml_parts_from_path,
    _sanitize_mapping_values,
    extract_placeholders_from_doc,
    safe_replace_placeholders,
)
from docx import Document as DocxFactory

MESES_PTBR = {
    1: "janeiro", 2: "fevereiro", 3: "março", 4: "abril", 5: "maio", 6: "junho",
    7: "julho", 8: "agosto", 9: "setembro", 10: "outubro", 11: "novembro", 12: "dezembro",
}


def get_justificativa_template_path() -> Path:
    """Retorna o caminho do modelo modelo_Justificativa.docx."""
    base_dir = Path(settings.BASE_DIR) / "viagens" / "documents"
    path = base_dir / "modelo_Justificativa.docx"
    if not path.exists():
        raise FileNotFoundError(
            f"Modelo de justificativa não encontrado: {path}"
        )
    return path


def _is_clean_placeholder_key(key: str) -> bool:
    """Ignora chaves que parecem XML ou lixo (evita exibir conteúdo interno do DOCX no formulário)."""
    if not key or not key.strip():
        return False
    k = key.strip()
    if "<" in k or ">" in k or "w:" in k or k.startswith("{"):
        return False
    if re.search(r"^[\s\wáéíóúâêôãõç\-_]+$", k, re.IGNORECASE) is None:
        return False
    return True


def get_justificativa_placeholders() -> list[str]:
    """
    Extrai do modelo DOCX a lista de placeholders (ordem alfabética).
    Usa o texto dos parágrafos (python-docx), não o XML bruto, para evitar
    exibir fragmentos OOXML no formulário.
    """
    template_path = get_justificativa_template_path()
    doc = DocxFactory(str(template_path))
    counts = extract_placeholders_from_doc(doc)
    keys = [k for k in counts if _is_clean_placeholder_key(k)]
    if keys:
        return sorted(keys)
    # Fallback: extração pelo XML (pode incluir lixo; filtramos depois)
    parts = _iter_docx_xml_parts_from_path(str(template_path))
    raw = _extract_placeholders(parts)
    keys = [k for k in raw if _is_clean_placeholder_key(k)]
    return sorted(keys)


def build_justificativa_docx_bytes(mapping: dict[str, str]) -> BytesIO:
    """
    Gera o DOCX da justificativa substituindo os placeholders do modelo.
    mapping: dicionário chave -> valor (chave = nome do placeholder sem chaves).
    Placeholders não informados são preenchidos com string vazia.
    """
    template_path = get_justificativa_template_path()
    doc = DocxFactory(str(template_path))

    template_placeholders = set(get_justificativa_placeholders())
    for key in template_placeholders:
        if key not in mapping:
            mapping[key] = ""
    mapping = _sanitize_mapping_values(mapping)
    safe_replace_placeholders(doc, mapping)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def replace_oficio_numero_no_texto(texto: str, numero_ano: str) -> str:
    """
    Substitui no texto os placeholders do número do ofício por numero_ano (ex: 123/2026).
    Aceita: "X/ANO", "Ofício nº X/ANO", "Ofício X/ANO" (case insensitive).
    """
    if not (numero_ano or "").strip():
        return texto
    if not texto:
        return ""
    import re
    valor = numero_ano.strip()
    # Ofício nº X/ANO ou Ofício n.º X/ANO
    texto = re.sub(r"Of[ií]cio\s*n[º°.]?\s*X\s*/\s*ANO", f"Ofício nº {valor}", texto, flags=re.IGNORECASE)
    # Ofício X/ANO
    texto = re.sub(r"Of[ií]cio\s+X\s*/\s*ANO", f"Ofício {valor}", texto, flags=re.IGNORECASE)
    # X/ANO sozinho (palavra)
    texto = re.sub(r"\bX\s*/\s*ANO\b", valor, texto, flags=re.IGNORECASE)
    return texto


def _format_data_extenso(d: date) -> str:
    """Data por extenso em português (ex: 3 de março de 2026)."""
    mes = MESES_PTBR.get(d.month, str(d.month))
    return f"{d.day} de {mes} de {d.year}"


def build_justificativa_context_from_config(
    *,
    assinante_nome: str = "",
    assinante_cargo: str = "",
    justificativa_texto: str = "",
    data_doc: date | None = None,
) -> dict[str, str]:
    """
    Monta o dicionário de contexto para o template da justificativa a partir
    da configuração do ofício (Configurações), data de hoje e assinante/texto.
    """
    from django.utils import timezone
    from viagens.services.oficio_config import get_oficio_config
    from viagens.services.text import title_case_pt

    cfg = get_oficio_config()
    hoje = data_doc or timezone.localdate()

    unidade = (getattr(cfg, "unidade_nome", "") or "").strip()
    unidade_rodape = title_case_pt(unidade)
    divisao = (getattr(cfg, "origem_nome", "") or "").strip()
    endereco = _build_endereco_formatado(cfg)
    telefone = (getattr(cfg, "telefone", "") or "").strip()
    email = (getattr(cfg, "email", "") or "").strip()

    sede = ""
    sede_cidade = getattr(cfg, "sede_cidade_default", None)
    if sede_cidade:
        nome = getattr(sede_cidade, "nome", "") or ""
        estado = getattr(sede_cidade, "estado", None)
        sigla = getattr(estado, "sigla", "") or "" if estado else ""
        sede = f"{nome}/{sigla}" if nome and sigla else nome or sigla

    return {
        "sede": sede,
        "data_extenso": _format_data_extenso(hoje),
        "justificativa": (justificativa_texto or "").strip(),
        "assinante_justificativa": (assinante_nome or "").strip(),
        "cargo_assinante_justificativa": (assinante_cargo or "").strip(),
        "divisao": divisao,
        "unidade": unidade,
        "unidade_rodape": unidade_rodape,
        "endereco": endereco,
        "telefone": telefone,
        "email": email,
    }
