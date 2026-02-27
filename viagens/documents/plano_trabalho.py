from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.utils import timezone
from django.utils.dateparse import parse_time
from docx import Document as DocxFactory

from viagens.diarias import PeriodMarker
from viagens.models import (
    Oficio,
    PlanoTrabalho,
    PlanoTrabalhoAtividade,
    PlanoTrabalhoLocalAtuacao,
    PlanoTrabalhoMeta,
    PlanoTrabalhoRecurso,
    Trecho,
    get_next_plano_num,
)
from viagens.services.oficio_config import get_oficio_config
from viagens.services.diarias_unified import (
    calculate_diarias_from_markers,
    derive_financeiro_diarias,
    parse_decimal_br,
)
from viagens.documents.document import (
    _find_unresolved_placeholders,
    apply_document_text_hygiene,
    extract_placeholders_from_doc,
    remove_optional_placeholder_paragraphs,
    safe_replace_placeholders,
)
from viagens.services.plano_trabalho import (
    ATIVIDADES_ORDEM_FIXA,
    META_POR_ATIVIDADE,
    PLANO_DOCX_OPTIONAL_PLACEHOLDERS,
    build_plano_placeholders,
    resolve_plano_docx_placeholders,
    validate_required_placeholders,
)

PLANO_TEMPLATE_FILENAME = "modelo_plano_de_trabalho.docx"


def _parse_decimal(value) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        normalized = str(value).strip().replace(".", "").replace(",", ".")
        return Decimal(normalized)
    except (InvalidOperation, TypeError, ValueError):
        return None


def _resolve_periodo(oficio: Oficio, trechos: list[Trecho]) -> tuple[date, date]:
    datas_inicio = [trecho.saida_data for trecho in trechos if trecho.saida_data]
    datas_fim = [
        trecho.chegada_data or trecho.saida_data
        for trecho in trechos
        if trecho.chegada_data or trecho.saida_data
    ]
    if oficio.retorno_saida_data:
        datas_fim.append(oficio.retorno_saida_data)
    if oficio.retorno_chegada_data:
        datas_fim.append(oficio.retorno_chegada_data)

    hoje = timezone.localdate()
    data_inicio = min(datas_inicio) if datas_inicio else hoje
    data_fim = max(datas_fim) if datas_fim else data_inicio
    if data_fim < data_inicio:
        data_fim = data_inicio
    return data_inicio, data_fim


def _resolve_destino(oficio: Oficio, trechos: list[Trecho]) -> str:
    labels: list[str] = []
    seen: set[str] = set()
    for trecho in trechos:
        if trecho.destino_cidade and trecho.destino_estado:
            label = f"{trecho.destino_cidade.nome}/{trecho.destino_estado.sigla}"
        elif trecho.destino_cidade:
            label = trecho.destino_cidade.nome
        elif trecho.destino_estado:
            label = trecho.destino_estado.sigla
        else:
            label = ""
        if label and label not in seen:
            seen.add(label)
            labels.append(label)
    if not labels:
        if oficio.cidade_destino and oficio.estado_destino:
            return f"{oficio.cidade_destino.nome}/{oficio.estado_destino.sigla}"
        if oficio.cidade_destino:
            return oficio.cidade_destino.nome
        return ""
    if len(labels) == 1:
        return labels[0]
    return ", ".join(labels[:-1]) + f" e {labels[-1]}"


def _resolve_local(oficio: Oficio, trechos: list[Trecho]) -> str:
    if oficio.cidade_sede and oficio.estado_sede:
        return f"{oficio.cidade_sede.nome}/{oficio.estado_sede.sigla}"
    if oficio.cidade_sede:
        return oficio.cidade_sede.nome
    for trecho in trechos:
        if trecho.destino_cidade and trecho.destino_estado:
            return f"{trecho.destino_cidade.nome}/{trecho.destino_estado.sigla}"
        if trecho.destino_cidade:
            return trecho.destino_cidade.nome
    cfg = get_oficio_config()
    if cfg.sede_cidade_default and cfg.sede_cidade_default.estado:
        return f"{cfg.sede_cidade_default.nome}/{cfg.sede_cidade_default.estado.sigla}"
    return "Curitiba/PR"


def _ensure_plano_trabalho(oficio: Oficio, trechos: list[Trecho]) -> PlanoTrabalho:
    try:
        return oficio.plano_trabalho
    except PlanoTrabalho.DoesNotExist:
        pass

    cfg = get_oficio_config()
    ano = int(oficio.ano or timezone.localdate().year)
    data_inicio, data_fim = _resolve_periodo(oficio, trechos)
    destino = _resolve_destino(oficio, trechos)
    qtd_servidores = int(oficio.viajantes.count() or 0)
    valor_total_legacy = _parse_decimal(oficio.valor_diarias)

    plano = PlanoTrabalho.objects.create(
        oficio=oficio,
        numero=get_next_plano_num(ano),
        ano=ano,
        sigla_unidade="ASCOM",
        programa_projeto="PCPR na Comunidade",
        destino=destino,
        solicitante="Demanda institucional",
        local=_resolve_local(oficio, trechos),
        data_inicio=data_inicio,
        data_fim=data_fim,
        horario_atendimento="das 09h as 17h",
        efetivo_formatado=f"{qtd_servidores} servidores.",
        efetivo_por_dia=qtd_servidores,
        quantidade_servidores=qtd_servidores,
        composicao_diarias=(oficio.quantidade_diarias or "").strip() or "1 x 100%",
        valor_total_calculado=valor_total_legacy,
        valor_unitario=valor_total_legacy,
        coordenador_plano=getattr(cfg, "assinante", None),
        coordenador_nome=(cfg.assinante.nome if getattr(cfg, "assinante", None) else ""),
        coordenador_cargo=(cfg.assinante.cargo if getattr(cfg, "assinante", None) else ""),
    )
    atividade_padrao = ATIVIDADES_ORDEM_FIXA[0]
    meta_padrao = META_POR_ATIVIDADE[atividade_padrao]
    PlanoTrabalhoMeta.objects.create(
        plano=plano,
        ordem=1,
        descricao=meta_padrao,
    )
    PlanoTrabalhoAtividade.objects.create(
        plano=plano,
        ordem=1,
        descricao=atividade_padrao,
    )
    PlanoTrabalhoRecurso.objects.create(
        plano=plano,
        ordem=1,
        descricao="Unidade movel da PCPR.",
    )
    PlanoTrabalhoLocalAtuacao.objects.create(
        plano=plano,
        ordem=1,
        data=data_inicio,
        local=destino or _resolve_local(oficio, trechos),
    )
    return plano


def _build_period_markers_from_trechos(trechos: list[Trecho]) -> list[PeriodMarker]:
    markers: list[PeriodMarker] = []
    for trecho in trechos:
        if not trecho.saida_data:
            continue
        saida_hora = trecho.saida_hora or time.min
        saida_dt = datetime.combine(trecho.saida_data, saida_hora)
        destino_uf = ""
        if trecho.destino_estado:
            destino_uf = trecho.destino_estado.sigla
        elif trecho.destino_cidade and trecho.destino_cidade.estado:
            destino_uf = trecho.destino_cidade.estado.sigla
        markers.append(
            PeriodMarker(
                saida=saida_dt,
                destino_cidade=trecho.destino_cidade.nome if trecho.destino_cidade else "",
                destino_uf=destino_uf,
            )
        )
    return markers


def _sync_plano_financeiro_from_diarias(
    plano: PlanoTrabalho,
    oficio: Oficio,
    trechos: list[Trecho],
) -> None:
    total_servidores = int(plano.quantidade_servidores or 0)
    if total_servidores < 1:
        raise ValueError(
            "Plano de trabalho incompleto. Preencha o efetivo para calcular as diarias."
        )
    markers = _build_period_markers_from_trechos(trechos)
    if not markers:
        raise ValueError(
            "Plano de trabalho incompleto. Informe trechos validos para calcular as diarias."
        )
    chegada_data = oficio.retorno_chegada_data or plano.data_fim
    chegada_hora = oficio.retorno_chegada_hora or parse_time("18:00") or time.min
    chegada_final = datetime.combine(chegada_data, chegada_hora)
    resultado = calculate_diarias_from_markers(
        markers=markers,
        chegada_final_sede=chegada_final,
        total_servidores=total_servidores,
    )
    financeiro = derive_financeiro_diarias(resultado)

    composicao = financeiro.get("diarias_por_servidor", "")
    valor_unitario = (
        parse_decimal_br(financeiro.get("valor_unitario"))
        or parse_decimal_br(financeiro.get("valor_por_servidor"))
    )
    valor_total = parse_decimal_br(financeiro.get("total_geral"))
    if not valor_total:
        raise ValueError(
            "Plano de trabalho incompleto. Nao foi possivel derivar o valor total das diarias."
        )

    update_fields: list[str] = []
    if composicao and plano.composicao_diarias != composicao:
        plano.composicao_diarias = composicao
        update_fields.append("composicao_diarias")
    if valor_unitario is not None and plano.valor_unitario != valor_unitario:
        plano.valor_unitario = valor_unitario
        update_fields.append("valor_unitario")
    if plano.valor_total_calculado != valor_total:
        plano.valor_total_calculado = valor_total
        update_fields.append("valor_total_calculado")

    if update_fields:
        plano.save(update_fields=update_fields + ["updated_at"])


def _resolve_plano_template_path() -> Path:
    return Path(settings.BASE_DIR) / "viagens" / "documents" / PLANO_TEMPLATE_FILENAME


def build_plano_trabalho_docx_bytes(oficio: Oficio) -> BytesIO:
    trechos = list(
        oficio.trechos.select_related(
            "origem_cidade",
            "origem_estado",
            "destino_cidade",
            "destino_estado",
        ).order_by("ordem", "id")
    )
    plano = _ensure_plano_trabalho(oficio, trechos)
    _sync_plano_financeiro_from_diarias(plano, oficio, trechos)
    cfg = get_oficio_config()
    placeholders = build_plano_placeholders(plano, oficio, cfg)
    missing = validate_required_placeholders(placeholders)
    if missing:
        raise ValueError(
            "Plano de trabalho incompleto. Campos obrigatorios ausentes: "
            + ", ".join(missing)
        )
    template_path = _resolve_plano_template_path()
    if not template_path.exists():
        raise FileNotFoundError(
            f"Template do plano nao encontrado: {template_path}"
        )
    doc = DocxFactory(str(template_path))
    template_placeholders = extract_placeholders_from_doc(doc)
    resolved_placeholders, missing_required_template = resolve_plano_docx_placeholders(
        template_placeholders.keys(),
        placeholders,
    )
    if missing_required_template:
        raise ValueError(
            "Plano de trabalho incompleto. Campos obrigatorios ausentes no template: "
            + ", ".join(sorted(missing_required_template))
        )

    remove_optional_placeholder_paragraphs(
        doc,
        resolved_mapping=resolved_placeholders,
        optional_placeholders=PLANO_DOCX_OPTIONAL_PLACEHOLDERS,
    )
    safe_replace_placeholders(doc, resolved_placeholders)
    apply_document_text_hygiene(doc)

    buf = BytesIO()
    doc.save(buf)
    leftovers = _find_unresolved_placeholders(buf.getvalue())
    if leftovers:
        raise ValueError(
            "Placeholders nao substituidos no plano DOCX: " + ", ".join(sorted(leftovers))
        )
    buf.seek(0)
    return buf
